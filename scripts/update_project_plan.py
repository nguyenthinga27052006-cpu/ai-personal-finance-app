# -*- coding: utf-8 -*-
"""
Script to update 01_GenAI_SoftwareDevelopment_project-plan.docx with the complete
9-week implementation plan for the 4 team members of class CNTT K23G.
"""
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import sys

def update_docx(docx_path):
    doc = docx.Document(docx_path)

    # 1. Update Header Paragraphs
    # Clear existing paragraphs and recreate them cleanly
    p0 = doc.paragraphs[0]
    p0.text = ""
    r0 = p0.add_run("KẾ HOẠCH THỰC HIỆN – PHÁT TRIỂN ỨNG DỤNG CHUYÊN SÂU")
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(14)
    r0.font.bold = True
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p1 = doc.paragraphs[1]
    p1.text = ""
    r1 = p1.add_run("Nhóm 02 – Thành viên nhóm (Lớp CNTT K23G – 04 sinh viên)")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.italic = True
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p2 = doc.paragraphs[2]
    p2.text = ""
    r2_b = p2.add_run("• Phạm Tiến Hà (Trưởng nhóm): ")
    r2_b.font.bold = True
    r2_b.font.name = "Times New Roman"
    r2_b.font.size = Pt(11)
    r2_t = p2.add_run("AI thông minh chatbot RAG, OCR, Backend, Quản trị hệ thống, kiểm soát tiến độ chung, DevOps")
    r2_t.font.name = "Times New Roman"
    r2_t.font.size = Pt(11)

    p3 = doc.paragraphs[3]
    p3.text = ""
    r3_b = p3.add_run("• Lưu Anh Đức: ")
    r3_b.font.bold = True
    r3_b.font.name = "Times New Roman"
    r3_b.font.size = Pt(11)
    r3_t = p3.add_run("Phân tích thiết kế hệ thống, Kiểm thử (Tester / QA-QC), Support cho các thành viên khác")
    r3_t.font.name = "Times New Roman"
    r3_t.font.size = Pt(11)

    p4 = doc.paragraphs[4]
    p4.text = ""
    r4_b = p4.add_run("• Lạc Mạnh Tuấn: ")
    r4_b.font.bold = True
    r4_b.font.name = "Times New Roman"
    r4_b.font.size = Pt(11)
    r4_t = p4.add_run("Phân tích trải nghiệm người dùng (UI/UX), Lập trình toàn diện giao diện Frontend, Thiết kế kiến trúc hệ thống, Thuyết trình chính")
    r4_t.font.name = "Times New Roman"
    r4_t.font.size = Pt(11)

    p5 = doc.paragraphs[5]
    p5.text = ""
    r5_b = p5.add_run("• Phạm Ngô Hoàng: ")
    r5_b.font.bold = True
    r5_b.font.name = "Times New Roman"
    r5_b.font.size = Pt(11)
    r5_t = p5.add_run("Quản lý tài liệu, Phân tích nghiệp vụ, Chuẩn bị dữ liệu kiểm thử mẫu và các công việc còn lại")
    r5_t.font.name = "Times New Roman"
    r5_t.font.size = Pt(11)

    p6 = doc.paragraphs[6]
    p6.text = ""
    r6_b = p6.add_run("Tên ứng dụng: ")
    r6_b.font.bold = True
    r6_b.font.name = "Times New Roman"
    r6_b.font.size = Pt(12)
    r6_t = p6.add_run("Xây dựng hệ thống Quản lý chi tiêu cá nhân tích hợp AI (AI Personal Finance Assistant)")
    r6_t.font.bold = True
    r6_t.font.name = "Times New Roman"
    r6_t.font.size = Pt(12)

    p7 = doc.paragraphs[7]
    p7.text = ""
    r7_b = p7.add_run("Thời gian thực hiện: ")
    r7_b.font.bold = True
    r7_b.font.name = "Times New Roman"
    r7_b.font.size = Pt(11.5)
    r7_t = p7.add_run("Từ 27/07/2026 đến 27/09/2026 (9 tuần)")
    r7_t.font.name = "Times New Roman"
    r7_t.font.size = Pt(11.5)

    p8 = doc.paragraphs[8]
    p8.text = ""
    r8 = p8.add_run("Kế hoạch chi tiết")
    r8.font.bold = True
    r8.font.name = "Times New Roman"
    r8.font.size = Pt(12.5)

    # 2. Define the 46 tasks for the 46 table rows (Rows 1 to 46 in Table 0)
    tasks_data = [
        # Tuần 01: 27/07/2026 - 02/08/2026 (5 tasks: row 1 -> 5)
        (
            "Khảo sát thực trạng quản lý tài chính cá nhân, phân tích bài toán và thu thập yêu cầu người dùng mục tiêu; lập kế hoạch quản lý tài liệu dự án",
            "Phạm Ngô Hoàng",
            "Hoàng phụ trách chính quản lý tài liệu và khảo sát yêu cầu"
        ),
        (
            "Xác định phạm vi đề tài, thiết lập cấu trúc phân rã công việc (WBS), thống nhất các mốc bàn giao và kiểm soát tiến độ chung toàn dự án",
            "Phạm Tiến Hà",
            "Hà phụ trách chính kiểm soát tiến độ và điều phối chung"
        ),
        (
            "Khảo sát các giải pháp AI tiên tiến cho tài chính: mô hình ngôn ngữ lớn (Gemini Flash), kỹ thuật RAG, trích xuất dữ liệu hóa đơn qua OCR Vision",
            "Phạm Tiến Hà",
            "Hà nghiên cứu phân hệ AI thông minh, RAG và OCR"
        ),
        (
            "Khảo sát hành vi người dùng, đánh giá trải nghiệm UI/UX các ứng dụng tài chính hàng đầu (Money Lover, Spendee) và định hướng công nghệ Frontend",
            "Lạc Mạnh Tuấn",
            "Tuấn phụ trách phân tích trải nghiệm người dùng (UX)"
        ),
        (
            "Khảo sát quy trình nghiệp vụ tài chính, đối soát dòng tiền và phân tích sơ bộ các quy chuẩn hệ thống; hỗ trợ các thành viên định hình bài toán",
            "Lưu Anh Đức",
            "Đức phân tích sơ bộ hệ thống và support nhóm"
        ),

        # Tuần 02: 03/08/2026 - 09/08/2026 (5 tasks: row 6 -> 10)
        (
            "Soạn thảo Tài liệu đặc tả yêu cầu phần mềm (SRS): mô tả chức năng User/Admin/AI, yêu cầu phi chức năng và tiêu chí nghiệm thu",
            "Phạm Ngô Hoàng",
            "Hoàng phụ trách hoàn thiện tài liệu SRS chi tiết"
        ),
        (
            "Phân tích thiết kế hệ thống mức quan niệm: xây dựng biểu đồ Use Case tổng thể và đặc tả kịch bản Use Case cho các phân hệ chức năng",
            "Lưu Anh Đức",
            "Đức phụ trách phân tích thiết kế Use Case hệ thống"
        ),
        (
            "Xây dựng luồng trải nghiệm người dùng (User Journey Map), phác thảo Wireframe các màn hình chính và cấu trúc điều hướng ứng dụng",
            "Lạc Mạnh Tuấn",
            "Tuấn phụ trách thiết kế trải nghiệm người dùng và Wireframe"
        ),
        (
            "Nghiên cứu kỹ thuật nhúng Vector (Embeddings), giải pháp pgvector trên PostgreSQL và thiết kế sơ bộ pipeline RAG cùng OCR bóc tách văn bản",
            "Phạm Tiến Hà",
            "Hà phụ trách kỹ thuật RAG, Vector Search và OCR"
        ),
        (
            "Thiết lập môi trường quản trị mã nguồn (Git branching, conventions), cấu hình base template Backend FastAPI và điều phối mục tiêu Sprint",
            "Phạm Tiến Hà",
            "Hà quản trị hệ thống và kiểm soát tiến độ kỹ thuật"
        ),

        # Tuần 03: 10/08/2026 - 16/08/2026 (5 tasks: row 11 -> 15)
        (
            "Thiết kế kiến trúc tổng thể hệ thống (System Architecture Diagram): mô hình phân tầng Client - API Gateway - Backend Services - AI Engine - Database",
            "Lạc Mạnh Tuấn",
            "Tuấn phụ trách chính thiết kế kiến trúc hệ thống"
        ),
        (
            "Thiết kế cấu trúc cơ sở dữ liệu quan hệ (ERD): bảng người dùng, ví/tài khoản, giao dịch, ngân sách, mục tiêu và bảng vector embeddings",
            "Phạm Tiến Hà",
            "Hà phụ trách thiết kế CSDL Backend và chỉ mục pgvector"
        ),
        (
            "Phân tích luồng dữ liệu (DFD), xây dựng biểu đồ tuần tự (Sequence Diagram) và biểu đồ hoạt động (Activity Diagram) cho luồng giao dịch tài chính",
            "Lưu Anh Đức",
            "Đức phân tích thiết kế luồng xử lý và dữ liệu hệ thống"
        ),
        (
            "Thiết kế chi tiết kiến trúc phân hệ AI: luồng xử lý RAG truy xuất ngữ cảnh tài chính, gateway kết nối Gemini LLM và cơ chế Function Calling",
            "Phạm Tiến Hà",
            "Hà thiết kế kiến trúc AI thông minh, Chatbot và OCR"
        ),
        (
            "Soạn thảo Tài liệu thiết kế hệ thống (SDD), tổng hợp sơ đồ kỹ thuật từ các thành viên và quản lý cấu trúc phiên bản tài liệu",
            "Phạm Ngô Hoàng",
            "Hoàng quản lý và tổng hợp tài liệu thiết kế hệ thống"
        ),

        # Tuần 04: 17/08/2026 - 23/08/2026 (5 tasks: row 16 -> 20)
        (
            "Thiết kế giao diện chi tiết (Figma/Prototype chuẩn Material 3, Dark/Light mode): Dashboard, Sổ thu chi, Ngân sách, Chatbot AI và Camera OCR",
            "Lạc Mạnh Tuấn",
            "Tuấn thiết kế toàn diện giao diện và trải nghiệm người dùng"
        ),
        (
            "Thiết kế chi tiết các RESTful API endpoints Backend (Auth, Accounts, Transactions, Budgets, Analytics, AI) và cơ chế phân quyền RBAC/JWT",
            "Phạm Tiến Hà",
            "Hà thiết kế chi tiết Backend API và bảo mật hệ thống"
        ),
        (
            "Xây dựng Kế hoạch kiểm thử tổng thể (Master Test Plan), thiết kế ma trận kiểm thử (Traceability Matrix) và xây dựng bộ Test Cases chi tiết",
            "Lưu Anh Đức",
            "Đức phụ trách chính kế hoạch kiểm thử và bộ Test Cases"
        ),
        (
            "Thiết kế cấu trúc Prompting chuyên sâu cho Trợ lý tài chính, quy tắc chống ảo giác (Guardrails) và bộ lọc dữ liệu cá nhân nhạy cảm (PII Masking)",
            "Phạm Tiến Hà",
            "Hà thiết kế giải pháp an toàn AI và Prompt Engineering"
        ),
        (
            "Xây dựng tài liệu đặc tả API chuẩn OpenAPI/Swagger và tổng hợp các quy định bảo vệ dữ liệu cá nhân & an toàn thông tin tài chính",
            "Phạm Ngô Hoàng",
            "Hoàng quản lý tài liệu API và các quy định pháp lý/bảo mật"
        ),

        # Tuần 05: 24/08/2026 - 30/08/2026 (6 tasks: row 21 -> 26)
        (
            "Khởi tạo hạ tầng DevOps: thiết lập Docker Compose (PostgreSQL 18 với extension pgvector, Redis Cache), cấu hình biến môi trường và logging",
            "Phạm Tiến Hà",
            "Hà phụ trách hạ tầng DevOps và quản trị hệ thống"
        ),
        (
            "Lập trình Backend phân hệ Xác thực người dùng (Auth Service): Đăng ký, Đăng nhập, mã hóa mật khẩu bcrypt, cấp phát JWT và quản lý Session",
            "Phạm Tiến Hà",
            "Hà phụ trách lập trình Backend Core API"
        ),
        (
            "Lập trình nền tảng ứng dụng Frontend Flutter: cấu hình điều hướng (GoRouter), quản lý trạng thái (Bloc/Cubit), hệ thống Theme và màn hình Auth",
            "Lạc Mạnh Tuấn",
            "Tuấn lập trình toàn diện giao diện Frontend"
        ),
        (
            "Thiết lập môi trường kiểm thử tự động (Pytest), viết bộ Unit Test cho module Auth và hỗ trợ gỡ lỗi tích hợp API cho các thành viên",
            "Lưu Anh Đức",
            "Đức thực thi kiểm thử và support kỹ thuật cho nhóm"
        ),
        (
            "Lập trình Gateway kết nối mô hình Gemini AI, thiết lập cấu trúc lưu trữ Vector Embeddings và kiểm thử hàm trích xuất giao dịch từ văn bản thô",
            "Phạm Tiến Hà",
            "Hà lập trình AI Gateway và kết nối mô hình thông minh"
        ),
        (
            "Cập nhật tiến độ dự án vào tài liệu kỹ thuật; chuẩn bị bộ dữ liệu mẫu (Seeding data: danh mục thu chi chuẩn Việt Nam, người dùng thử nghiệm)",
            "Phạm Ngô Hoàng",
            "Hoàng quản lý tài liệu và chuẩn bị dữ liệu seeding ban đầu"
        ),

        # Tuần 06: 31/08/2026 - 06/09/2026 (5 tasks: row 27 -> 31)
        (
            "Lập trình Backend phân hệ Quản lý tài khoản/ví (Tiền mặt, Ngân hàng, Thẻ tín dụng) và phân hệ Giao dịch tài chính (Thu, Chi, Chuyển khoản nội bộ)",
            "Phạm Tiến Hà",
            "Hà lập trình Backend logic kế toán và quản lý giao dịch"
        ),
        (
            "Lập trình Backend phân hệ Hạn mức ngân sách (Budget Alerts theo dõi vượt ngưỡng), Mục tiêu tiết kiệm (Savings Goals) và API Dashboard thống kê",
            "Phạm Tiến Hà",
            "Hà hoàn thiện các phân hệ tài chính Backend cốt lõi"
        ),
        (
            "Lập trình toàn diện giao diện Frontend: Dashboard tài chính (Biểu đồ tròn tỷ trọng, Biểu đồ cột dòng tiền), Sổ giao dịch và Màn hình tạo giao dịch",
            "Lạc Mạnh Tuấn",
            "Tuấn lập trình giao diện Dashboard, Biểu đồ và Giao dịch"
        ),
        (
            "Thực thi kiểm thử chức năng (Functional Testing) cho phân hệ Quản lý ví, Giao dịch và Ngân sách; ghi nhận bug log và xác minh tính toàn vẹn số dư",
            "Lưu Anh Đức",
            "Đức làm tester kiểm thử nghiệp vụ và support các thành viên"
        ),
        (
            "Xây dựng bộ dữ liệu thử nghiệm chi tiết (100+ giao dịch thu chi mẫu đa dạng ngữ cảnh); rà soát và cập nhật tài liệu kỹ thuật các phân hệ đã xong",
            "Phạm Ngô Hoàng",
            "Hoàng quản lý tài liệu và phụ trách bộ dữ liệu kiểm thử mẫu"
        ),

        # Tuần 07: 07/09/2026 - 13/09/2026 (5 tasks: row 32 -> 36)
        (
            "Xây dựng pipeline RAG hoàn chỉnh: vector hóa lịch sử tài chính với pgvector, truy xuất ngữ cảnh liên quan và tối ưu prompt ngữ cảnh người dùng",
            "Phạm Tiến Hà",
            "Hà phụ trách chính phân hệ Chatbot RAG và Vector Search"
        ),
        (
            "Xây dựng module OCR Vision AI: bóc tách tự động thông tin từ ảnh chụp hóa đơn (Tên cửa hàng, Ngày giờ, Tổng tiền, Chi tiết món, Gợi ý danh mục)",
            "Phạm Tiến Hà",
            "Hà lập trình Backend module OCR trích xuất hóa đơn"
        ),
        (
            "Lập trình giao diện Trợ lý ảo Chatbot (hiệu ứng streaming typing, thẻ tóm tắt chi tiêu) và giao diện Quét hóa đơn OCR (chụp/chọn ảnh, preview, xác nhận)",
            "Lạc Mạnh Tuấn",
            "Tuấn lập trình giao diện Chatbot AI và Camera OCR Scanner"
        ),
        (
            "Thực hiện kiểm thử phân hệ AI & OCR: kiểm thử độ chính xác trích xuất hóa đơn, đo độ trễ LLM, kiểm thử an toàn prompt và hỗ trợ kết nối API AI",
            "Lưu Anh Đức",
            "Đức kiểm thử phân hệ AI, OCR và support tích hợp Client"
        ),
        (
            "Thu thập và phân loại 50+ ảnh hóa đơn thực tế làm tập dữ liệu mẫu cho OCR; soạn thảo kịch bản hội thoại mẫu phục vụ kiểm thử trợ lý ảo AI",
            "Phạm Ngô Hoàng",
            "Hoàng quản lý tài liệu, dữ liệu mẫu OCR và kịch bản AI"
        ),

        # Tuần 08: 14/09/2026 - 20/09/2026 (5 tasks: row 37 -> 41)
        (
            "Triển khai kiểm thử toàn diện: Kiểm thử tích hợp (Integration Test), Kiểm thử luồng người dùng (E2E User Journey), Kiểm thử bảo mật và tải hệ thống",
            "Lưu Anh Đức",
            "Đức phụ trách toàn diện công tác kiểm thử (Tester Lead)"
        ),
        (
            "Khắc phục các lỗi phát hiện qua kiểm thử; tối ưu hóa truy vấn CSDL PostgreSQL, cấu hình caching Redis và tối ưu tốc độ xử lý của AI/OCR",
            "Phạm Tiến Hà",
            "Hà tối ưu Backend, AI Engine và kiểm soát tiến độ chung"
        ),
        (
            "Tối ưu hóa hiệu năng giao diện Frontend (mượt mà 60 FPS, micro-animations, xử lý ngoại lệ mất mạng) và thiết kế slide thuyết trình đồ án",
            "Lạc Mạnh Tuấn",
            "Tuấn tối ưu trải nghiệm UI/UX và chuẩn bị slide thuyết trình"
        ),
        (
            "Đóng gói hoàn chỉnh hệ thống bằng Docker Compose chuẩn sản xuất (API, Worker, Postgres pgvector, Redis) và thiết lập kịch bản triển khai",
            "Phạm Tiến Hà",
            "Hà hoàn thiện đóng gói DevOps và quản trị hệ thống"
        ),
        (
            "Soạn thảo Báo cáo đồ án chính thức (các chương Tổng quan, Cơ sở lý thuyết, Phân tích thiết kế, Cài đặt & Đánh giá) và viết Tài liệu hướng dẫn sử dụng",
            "Phạm Ngô Hoàng",
            "Hoàng hoàn thiện tài liệu Báo cáo đồ án và Hướng dẫn sử dụng"
        ),

        # Tuần 09: 21/09/2026 - 27/09/2026 (5 tasks: row 42 -> 46)
        (
            "Triển khai chính thức hệ thống lên môi trường máy chủ thử nghiệm, kiểm tra tính sẵn sàng hạ tầng và thiết lập cơ chế sao lưu dữ liệu tự động",
            "Phạm Tiến Hà",
            "Hà phụ trách triển khai DevOps, vận hành và quản trị hệ thống"
        ),
        (
            "Kiểm tra lần cuối toàn bộ môi trường và dữ liệu sau triển khai (Smoke Test); chuẩn bị các phương án xử lý sự cố kỹ thuật dự phòng khi demo",
            "Lưu Anh Đức",
            "Đức kiểm thử xác nhận triển khai và support kỹ thuật demo"
        ),
        (
            "Hoàn thiện bài thuyết trình chuyên nghiệp, đảm nhận vai trò Thuyết trình chính trước Hội đồng chấm thi và trực tiếp điều hành phiên Demo sản phẩm",
            "Lạc Mạnh Tuấn",
            "Tuấn đảm nhiệm vai trò Thuyết trình chính và điều hành Demo"
        ),
        (
            "Rà soát tính nhất quán của toàn bộ tài liệu dự án, chuẩn bị bộ câu hỏi và kịch bản trả lời phản biện (Q&A); hoàn thiện hồ sơ nghiệm thu sản phẩm",
            "Phạm Ngô Hoàng",
            "Hoàng quản lý tài liệu, kịch bản phản biện và hồ sơ nghiệm thu"
        ),
        (
            "Tổng duyệt kịch bản demo (Đăng nhập -> Quét hóa đơn OCR -> Ghi nhận tự động -> Chat RAG hỏi đáp tài chính) và hoàn thành bảo vệ đồ án",
            "Phạm Tiến Hà,\nLưu Anh Đức,\nLạc Mạnh Tuấn,\nPhạm Ngô Hoàng",
            "Cả 4 thành viên cùng tham gia tổng duyệt và bảo vệ đồ án"
        )
    ]

    table = doc.tables[0]
    
    # Check that we have exactly 46 tasks matching rows 1 to 46
    assert len(tasks_data) == 46, f"Expected 46 tasks, got {len(tasks_data)}"
    assert len(table.rows) == 47, f"Expected 47 rows, got {len(table.rows)}"

    for idx, (job, member, note) in enumerate(tasks_data):
        row_idx = idx + 1
        row = table.rows[row_idx]
        
        # Col 1: Công việc
        cell1 = row.cells[1]
        cell1.text = ""
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.line_spacing = 1.15
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(job)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)

        # Col 2: Thành viên thực hiện
        cell2 = row.cells[2]
        cell2.text = ""
        p2 = cell2.paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(member)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)
        r2.font.bold = True

        # Col 3: Ghi chú
        cell3 = row.cells[3]
        cell3.text = ""
        p3 = cell3.paragraphs[0]
        p3.paragraph_format.line_spacing = 1.15
        p3.paragraph_format.space_before = Pt(2)
        p3.paragraph_format.space_after = Pt(2)
        p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r3 = p3.add_run(note)
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(10)
        r3.font.italic = True

    doc.save(docx_path)
    print(f"Successfully updated: {docx_path}")

if __name__ == "__main__":
    target = os.path.abspath("01_GenAI_SoftwareDevelopment_project-plan.docx")
    update_docx(target)
